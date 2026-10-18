"""Genere un petit corpus de demonstration (secteur transport) dans samples/."""
from pathlib import Path

import pandas as pd
import pymupdf as fitz
from docx import Document

out = Path("samples")
out.mkdir(exist_ok=True)

d = Document()
d.add_heading("Procedure de maintenance du portique de validation PV-300", level=1)
d.add_paragraph("Cette procedure s'applique aux portiques PV-300 installes dans les stations du reseau.")
d.add_heading("Securite", level=2)
d.add_paragraph("Avant toute ouverture du capot, couper l'alimentation 48 V au disjoncteur Q2 et attendre 5 minutes "
                "pour la decharge des condensateurs. Le port des gants isolants classe 0 est obligatoire.")
d.add_heading("Caracteristiques techniques", level=2)
t = d.add_table(rows=5, cols=2)
for i, (k, v) in enumerate([("Parametre", "Valeur"), ("Indice de protection", "IP54"),
                            ("Temperature de fonctionnement", "-20 C a +55 C"),
                            ("Debit maximal", "45 passagers par minute"), ("Interface", "RS485 / Ethernet")]):
    t.cell(i, 0).text, t.cell(i, 1).text = k, v
d.add_heading("Frequence de maintenance", level=2)
d.add_paragraph("Nettoyage des lecteurs sans contact : tous les 15 jours. Verification des capteurs de passage : "
                "tous les 3 mois. Remplacement des batteries de secours : tous les 2 ans.")
d.save(out / "procedure_PV300.docx")

pdf = fitz.open()
page = pdf.new_page()
page.insert_text((50, 60), "Reglement voyageurs - Ligne T2", fontsize=16)
y = 100
for line in ["Article 4 - Bagages : chaque voyageur peut transporter deux bagages a main.",
             "Les velos pliants sont acceptes a bord sans supplement.",
             "Article 7 - Retards : un retard superieur a 30 minutes ouvre droit",
             "a un avoir de 50 % du prix du billet, sur demande sous 60 jours.",
             "Article 9 - Objets trouves : ils sont conserves 3 mois au centre de Lyon Part-Dieu."]:
    page.insert_text((50, y), line, fontsize=11)
    y += 22
pdf.save(out / "reglement_T2.pdf")

pd.DataFrame({
    "Date": pd.to_datetime(["2026-09-01", "2026-09-02", "2026-09-02", "2026-09-05", "2026-09-08", "2026-09-09"]),
    "Ligne": ["T1", "T2", "T2", "M1", "T2", "M1"],
    "Type d'incident": ["Panne portique", "Signalisation", "Panne portique", "Malaise voyageur", "Signalisation",
                        "Panne portique"],
    "Retard (min)": [12, 35, 8, 20, 41, 5],
}).to_excel(out / "incidents_septembre.xlsx", index=False)
print("ok", sorted(p.name for p in out.iterdir()))
