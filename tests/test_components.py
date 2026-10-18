import pytest
from langchain_core.documents import Document

from app.graph.builder import _is_follow_up, format_context, mentions_schema, parse_route
from app.ingestion import _slug, split_documents
from app.loaders import StructuredDocumentLoader
from app.sql_backend import UnsafeSQLError, clean_sql, validate_sql


@pytest.mark.parametrize("sql", [
    "SELECT * FROM tabular.t",
    "with x as (select 1) select * from x;",
    'SELECT "date_update" FROM tabular.t',  # mot interdit en sous-chaine : autorise
])
def test_validate_sql_accepts_reads(sql):
    assert validate_sql(sql)


@pytest.mark.parametrize("sql", [
    "DROP TABLE tabular.t",
    "SELECT 1; DELETE FROM tabular.t",
    "UPDATE tabular.t SET a = 1",
    "SELECT pg_sleep(10)",
    "",
])
def test_validate_sql_rejects_writes(sql):
    with pytest.raises(UnsafeSQLError):
        validate_sql(sql)


def test_clean_sql_strips_markdown_and_preamble():
    assert clean_sql("Voici la requete :\n```sql\nSELECT 1\n```") == "SELECT 1"
    assert clean_sql("SELECT 1; SELECT 2;") == "SELECT 1"


@pytest.mark.parametrize("raw,has_tables,expected", [
    ("documents", True, "documents"),
    ("Database.", True, "database"),
    ("database", False, "documents"),
    ("chat", True, "chat"),
    ("je ne sais pas", True, "documents"),
])
def test_parse_route(raw, has_tables, expected):
    assert parse_route(raw, has_tables) == expected


def test_router_guard_vetoes_sql_without_analytic_intent():
    assert parse_route("database", True, "Quel est l'indice de protection du PV-300 ?") == "documents"
    assert parse_route("database", True, "Combien d'incidents par ligne ?") == "database"
    assert parse_route("database", True, "Retard cumulé sur la T2") == "database"


def test_router_guard_keeps_chat_for_short_messages_only():
    assert parse_route("chat", True, "Bonjour !") == "chat"
    assert parse_route("chat", True, "Quel est le delai pour demander un avoir sur la ligne T2 ?") == "documents"


def test_format_context_cites_page_only_when_known():
    ctx = format_context([
        Document(page_content="a", metadata={"source": "m.pdf", "page": 2}),
        Document(page_content="b", metadata={"source": "n.docx", "page": None}),
    ])
    assert "[m.pdf, page 2]" in ctx and "[n.docx]" in ctx


def test_slug():
    assert _slug("N° Série / Instrument") == "n_serie_instrument"


def test_docx_loader_and_splitter(tmp_path):
    from docx import Document as Docx

    path = tmp_path / "procedure.docx"
    d = Docx()
    d.add_heading("Maintenance", level=1)
    d.add_paragraph("Couper l'alimentation avant d'ouvrir le boitier. " * 80)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Parametre", "Valeur"
    t.cell(1, 0).text, t.cell(1, 1).text = "IP", "66"
    d.save(path)

    docs = StructuredDocumentLoader(path).load()
    assert docs and all(x.metadata["source"] == "procedure.docx" for x in docs)
    assert "66" in "\n".join(x.page_content for x in docs)

    chunks = split_documents(docs)
    assert len(chunks) > 1
    assert chunks[0].page_content.startswith("[procedure.docx]")


TABLE_INFO = """CREATE TABLE tabular.incidents_septembre (date TIMESTAMP, ligne TEXT, type_d_incident TEXT, retard_min BIGINT)
/* 3 rows from incidents_septembre table:
date\tligne\ttype_d_incident\tretard_min
2026-09-01\tT1\tPanne portique\t12 */"""


def test_mentions_schema():
    assert mentions_schema("Combien d'incidents par ligne ?", TABLE_INFO)
    assert mentions_schema("Retard cumulé sur la T2", TABLE_INFO)
    assert not mentions_schema("Combien de temps sont conservés les objets trouvés ?", TABLE_INFO)


def test_follow_up_detection():
    assert _is_follow_up("Et sa temperature ?")
    assert not _is_follow_up("Que faut-il faire avant d'ouvrir le capot ?")
